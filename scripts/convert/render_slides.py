"""Render the training slide decks for the on-site slide player (task M2-07).

Decks are played on the site only: no PDF, no download (user decision, 2026-10-26).

Usage:
    python scripts/convert/render_slides.py [deck-slug ...]

For each deck (all 11 by default):
  1. copies the .pptx to a temporary folder (the source is never opened for writing);
  2. hides every picture that docs/image-review.csv marks `drop` or `redraw`, matched by
     the same content hash pptx2mdx.py uses for its file names;
  3. exports each slide with PowerPoint (COM) at 1600x900 and saves it as WebP (<= 200 KB)
     to src/assets/slides/<deck>/NN.webp;
  4. writes src/data/slides/<deck>.json: title, source, credit, and each slide's title,
     text (Markdown) and speaker notes.

Skipped: the title slide, thank-you / questions slides, and diagram slides that the review
marks `drop` or `redraw` (the whole slide is the third-party figure).
Windows only (needs PowerPoint and pywin32).
"""
import csv
import hashlib
import io
import json
import os
import re
import shutil
import sys
import tempfile

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx2mdx import nice_title  # noqa: E402
from pptx2mdx import REVIEW, SKIP_TITLES, clean, find_title, is_title, table_md, text_frame_md, walk  # noqa: E402

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
SRC = os.path.join(os.path.dirname(ROOT), "Bhutan ART-TREES HFLD Readiness Training")
ASSETS = os.path.join(ROOT, "src", "assets", "slides")
DATA = os.path.join(ROOT, "src", "data", "slides")
MAX_BYTES = 200_000
SIZE = (1600, 900)

# deck slug -> (source path under SRC, day, lesson slug the deck belongs to)
DECKS = {
    "introduction-to-earth-engine": ("Day 1/D1 Introduction to GEE_ Features, Capabilities, and Access.pptx", 1, "day-1/introduction-to-earth-engine"),
    "code-editor-and-data-catalog": ("Day 1/D1 EXPLORING DATA ARCHIVE IN GEE.pptx", 1, "day-1/code-editor-and-data-catalog"),
    "image-collections": ("Day 1/D1 WORKING WITH IMAGE COLLECTIONS.pptx", 1, "day-1/image-collections"),
    "remote-sensing-fundamentals": ("Day 2/D2-Fundamentals_of_RS.pptx", 2, "day-2/remote-sensing-fundamentals"),
    "mosaics-and-composites": ("Day 2/D2-Creating Mosaics and Composites from Image Collections.pptx", 2, "day-2/mosaics-and-composites"),
    "import-and-export": ("Day 2/D2-Importing and Exporting Datasets with Google Earth Engine (GEE).pptx", 2, "day-2/import-and-export"),
    "random-forest": ("Day 3/D3 INTRO TO RANDOM FOREST.pptx", 3, "day-3/random-forest"),
    "rlcms-overview": ("Day 4/D4 RLCMS - GENERAL OVERVIEW AND TERMS TO KNOW.pptx", 4, "day-4/rlcms-overview"),
    "satellite-embeddings": ("Day 4/D4-Use of Google_s Satellite Embedding dataset.pptx", 4, "day-4/satellite-embeddings"),
    "sampling-design": ("Day 5/D5-Designing a Stratified Random Sampling Frame.pptx", 5, "day-5/sampling-design"),
    "accuracy-and-area": ("Day 5/D5-Accuracy Assessment and Area Adjustment.pptx", 5, "day-5/accuracy-and-area"),
}
# Titles for decks whose title slide is a photo or a generic course title.
TITLES = {
    "introduction-to-earth-engine": "Introduction to Google Earth Engine: Features, Capabilities and Access",
    "remote-sensing-fundamentals": "Fundamentals of Remote Sensing",
    "mosaics-and-composites": "Creating Mosaics and Composites from Image Collections",
    "import-and-export": "Importing and Exporting Datasets with Google Earth Engine",
    "satellite-embeddings": "Using Google's Satellite Embedding Dataset",
}
CREDIT = ("Adapted, with permission, from the Bhutan ART-TREES HFLD Readiness Training (May 2026) "
          "by Ate Poortinga and Samuel, Conservation International. Third-party pictures removed. CC BY 4.0.")


def review_for(deck):
    """(hidden picture hashes, hidden diagram slide numbers) for a deck."""
    hidden, slides = set(), set()
    with open(REVIEW, encoding="utf8", newline="") as f:
        for r in csv.DictReader(f):
            if r["deck"] != deck or r["decision"] not in ("drop", "redraw"):
                continue
            if r["kind"] == "diagram":
                slides.add(int(r["slide"]))
            else:
                hidden.add(re.match(r"s\d+-([0-9a-f]{10})\.", r["file"]).group(1))
    return hidden, slides


def all_shapes(shapes):
    for s in shapes:
        yield s
        if s.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from all_shapes(s.shapes)


def slide_text(shapes, title_shape):
    body = []
    for s in shapes:
        if s is title_shape or is_title(s):
            continue
        if s.has_text_frame and not re.search(r"[A-Za-z]{2}", s.text_frame.text):
            continue
        if getattr(s, "has_table", False) and s.has_table:
            body.append(table_md(s.table))
        elif s.has_text_frame:
            md = text_frame_md(s.text_frame)
            if md:
                body.append(md)
    return "\n\n".join(b for b in body if b)


def plan(path, deck):
    """Which slides to export, which shapes to hide, and the text of each slide."""
    hidden, drop_slides = review_for(deck)
    prs = Presentation(path)
    out = []
    for n, slide in enumerate(prs.slides, 1):
        shapes = walk(slide.shapes)
        title_shape = find_title(shapes)
        title = re.sub(r"\s+", " ", clean(title_shape.text_frame.text.split("\n")[0])) if title_shape is not None else ""
        if n == 1 or (title and SKIP_TITLES.match(title)) or n in drop_slides:
            continue
        hide = []
        for s in all_shapes(slide.shapes):
            if s.shape_type == MSO_SHAPE_TYPE.PICTURE or (s.is_placeholder and hasattr(s, "image")):
                try:
                    key = hashlib.sha1(s.image.blob).hexdigest()[:10]
                except Exception:
                    continue
                if key in hidden:
                    hide.append(s.shape_id)
        notes = ""
        if slide.has_notes_slide:
            notes = "\n\n".join(p.strip() for p in clean(slide.notes_slide.notes_text_frame.text).split("\n") if p.strip())
        text = slide_text(shapes, title_shape)
        if not title and not text and not hide and not notes and len(slide.shapes) <= 2:
            continue  # empty or photo-only closing slide
        out.append({"n": n, "title": nice_title(title) if title else "", "text": text, "notes": notes, "hide": hide})
    return out


def deck_title(path):
    """The title slide's title, as pptx2mdx.py chooses it (a generic course title gives way to the subtitle)."""
    t = find_title(walk(Presentation(path).slides[0].shapes))
    if t is None:
        return None
    lines = [clean(p.text) for p in t.text_frame.paragraphs if clean(p.text)]
    title = lines[0] if lines else None
    if title and len(lines) > 1 and re.fullmatch(r"(bhutan\s+)?(art-trees\s+)?hfld readiness training", title.lower()):
        title = " ".join(lines[1:]).split(":")[-1].strip()
    return nice_title(re.sub(r"\s+", " ", title)) if title else None


def hide_shapes(com_slide, ids):
    def visit(shapes):
        for i in range(1, shapes.Count + 1):
            sh = shapes.Item(i)
            if sh.Id in ids:
                sh.Visible = 0
            elif sh.Type == 6:  # msoGroup
                visit(sh.GroupItems)
    visit(com_slide.Shapes)


def to_webp(png, dest):
    im = Image.open(png).convert("RGB")
    if im.size != SIZE:
        im = im.resize(SIZE, Image.LANCZOS)
    for q in (80, 72, 64, 56, 48, 40):
        buf = io.BytesIO()
        im.save(buf, "WEBP", quality=q, method=6)
        if buf.tell() <= MAX_BYTES:
            break
    with open(dest, "wb") as f:
        f.write(buf.getvalue())
    return buf.tell()


def render(slug, app):
    src_rel, day, lesson = DECKS[slug]
    src = os.path.join(SRC, *src_rel.split("/"))
    deck = os.path.basename(src)
    tmp = tempfile.mkdtemp(prefix="slides-")
    try:
        copy = os.path.join(tmp, "deck.pptx")
        shutil.copyfile(src, copy)
        slides = plan(copy, deck)
        dest = os.path.join(ASSETS, slug)
        shutil.rmtree(dest, ignore_errors=True)
        os.makedirs(dest)
        pres = app.Presentations.Open(copy, False, False, False)  # read-write copy, no window
        total = 0
        try:
            for i, s in enumerate(slides, 1):
                com = pres.Slides(s["n"])
                hide_shapes(com, set(s["hide"]))
                png = os.path.join(tmp, f"{i:02d}.png")
                com.Export(png, "PNG", *SIZE)
                s["image"] = f"{i:02d}.webp"
                total += to_webp(png, os.path.join(dest, s["image"]))
        finally:
            pres.Close()
        title = TITLES.get(slug) or deck_title(copy) or slug.replace("-", " ").capitalize()
        os.makedirs(DATA, exist_ok=True)
        data = {
            "deck": slug, "title": title, "day": day, "lesson": lesson,
            "source": f"{src_rel} (Bhutan ART-TREES HFLD Readiness Training, May 2026)",
            "credit": CREDIT,
            "slides": [{"image": s["image"], "source_slide": s["n"], "title": s["title"], "text": s["text"],
                        "notes": s["notes"], "hidden_pictures": len(s["hide"])} for s in slides],
        }
        with open(os.path.join(DATA, slug + ".json"), "w", encoding="utf8", newline="\n") as f:
            json.dump(data, f, ensure_ascii=False, indent=1)
            f.write("\n")
        print(f"{slug}: {len(slides)} slides, {sum(len(s['hide']) for s in slides)} pictures hidden, {total / 1e6:.2f} MB")
        return total
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def main():
    import win32com.client  # pywin32
    slugs = sys.argv[1:] or list(DECKS)
    unknown = [s for s in slugs if s not in DECKS]
    if unknown:
        sys.exit(f"unknown deck(s): {', '.join(unknown)}; choose from {', '.join(DECKS)}")
    app = win32com.client.DispatchEx("PowerPoint.Application")
    try:
        total = sum(render(s, app) for s in slugs)
    finally:
        app.Quit()
    print(f"total {total / 1e6:.1f} MB")


if __name__ == "__main__":
    main()
