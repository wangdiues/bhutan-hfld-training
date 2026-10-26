"""Convert a training slide deck (.pptx) into a draft lesson MDX page (task M2-03).

Slides become web content, not PDFs (user decision, 2026-10-26).

Usage:
    python scripts/convert/pptx2mdx.py <deck.pptx> <lesson-slug> [--render-diagrams] [--install-images]

    <lesson-slug> is the page path under src/content/docs, e.g. day-1/introduction-to-earth-engine.

Writes:
    scripts/convert/out/<slug>.mdx         draft page (gitignored), one section per slide
    scripts/convert/out/<slug>/*.webp      extracted pictures (deduplicated) and rendered diagrams
    scripts/convert/out/<slug>/images.csv  one row per image, for the review manifest (task M2-04)

Per slide: the title becomes a "##" heading (consecutive slides with the same title are
merged), text boxes become paragraphs or lists in reading order, text set mostly in a
monospace font becomes a ```js block, tables become Markdown tables, and speaker notes
become prose. Slides made mostly of drawn shapes (diagrams) are rendered to an image with
PowerPoint (--render-diagrams, Windows only).

Images are linked only if docs/image-review.csv marks them `keep` (with alt text); any
other image is left as a TODO comment, so nothing unreviewed is published.
With --install-images, kept images are copied to src/assets/<slug>/.
The source .pptx is only read, never modified.
"""
import argparse
import csv
import hashlib
import io
import os
import re
import shutil
import sys

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx2mdx import MONO_FONTS, PRIVATE_USE, escape_prose, nice_title, quote_yaml  # noqa: E402

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
OUT = os.path.join(ROOT, "scripts", "convert", "out")
REVIEW = os.path.join(ROOT, "docs", "image-review.csv")

MAX_WIDTH = 1600  # px; larger pictures are scaled down
WEBP_QUALITY = 80
DIAGRAM_SHAPES = 12  # a slide with at least this many drawn shapes is treated as a diagram
SKIP_TITLES = re.compile(r"^(thank\s*you|thanks|questions\??|q\s*&\s*a)[.!]*$", re.I)
CODE_LINE = re.compile(r"^\s*(var |let |const |function\b|//|Map\.|Export\.|print\(|ee\.|return\b|\}|[\w.]+\s*=\s*.*;\s*$|.*\);\s*$)")
DRAWN = {MSO_SHAPE_TYPE.FREEFORM, MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.LINE}


def walk(shapes):
    """Shapes in reading order (top to bottom, then left to right), groups flattened."""
    flat = []
    for s in shapes:
        if s.shape_type == MSO_SHAPE_TYPE.GROUP:
            flat.extend(walk(s.shapes))
        else:
            flat.append(s)
    return sorted(flat, key=lambda s: ((s.top or 0) // 100000, s.left or 0))


def clean(text):
    return PRIVATE_USE.sub("", text).replace("\x0b", " ").replace("\u00a0", " ").strip()


def is_title(shape):
    return shape.is_placeholder and shape.placeholder_format.type in (
        PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE)


def max_font_pt(tf):
    sizes = [r.font.size.pt for p in tf.paragraphs for r in p.runs if r.font.size is not None]
    return max(sizes) if sizes else 0


def find_title(shapes):
    """The title placeholder, else the short text box with the largest font (designed slides)."""
    for s in shapes:
        if is_title(s) and s.has_text_frame and clean(s.text_frame.text):
            return s
    best, best_pt = None, 0
    for s in shapes:
        if not s.has_text_frame:
            continue
        text = clean(s.text_frame.text)
        first = next((clean(p.text) for p in s.text_frame.paragraphs if clean(p.text)), "")
        if not re.search(r"[A-Za-z]{2}", text) or len(first.split()) > 14 or mono_share(s.text_frame) >= 0.6:
            continue
        pt = max_font_pt(s.text_frame)
        if pt > best_pt:
            best, best_pt = s, pt
    return best if best_pt >= 20 else None


def mono_share(tf):
    mono = total = 0
    for p in tf.paragraphs:
        for r in p.runs:
            n = len(r.text.replace(" ", ""))
            total += n
            if r.font.name in MONO_FONTS:
                mono += n
    return mono / total if total else 0


def text_frame_md(tf):
    paras = [(p.level, clean("".join(r.text for r in p.runs) or p.text)) for p in tf.paragraphs]
    paras = [(lvl, t) for lvl, t in paras if t]
    if not paras:
        return ""
    code_like = sum(bool(CODE_LINE.match(t)) for _, t in paras)
    if mono_share(tf) >= 0.6 or (code_like >= 2 and code_like >= len(paras) / 2):
        # Code: mostly monospace, or mostly lines that look like JavaScript. Prose lines
        # mixed in become comments so the block stays runnable.
        if mono_share(tf) < 0.6:
            lines = [t if CODE_LINE.match(t) else f"// {t}" for _, t in paras]
            return "```js\n" + "\n".join(lines) + "\n```"
        code = "\n".join(PRIVATE_USE.sub("", p.text).replace("\x0b", "\n").rstrip() for p in tf.paragraphs)
        return f"```js\n{code.strip(chr(10))}\n```"
    if len(paras) == 1 and paras[0][0] == 0:
        return escape_prose(paras[0][1])
    return "\n".join(f"{'  ' * lvl}- {escape_prose(t)}" for lvl, t in paras)


def table_md(table):
    if len(table.rows) == 1:
        # A one-row table on a slide is a layout box, not data: emit its cells as text or code.
        cells = [clean(c.text) for c in table.rows[0].cells if clean(c.text)]
        out = []
        for c in cells:
            lines = c.split("\n")
            if sum(bool(CODE_LINE.match(l)) for l in lines) >= max(1, len(lines) / 2):
                out.append("```js\n" + PRIVATE_USE.sub("", c) + "\n```")
            else:
                out.append(escape_prose(c))
        return "\n\n".join(out)
    rows = [[escape_prose(clean(c.text)).replace("|", "\\|").replace("\n", " ") or " " for c in r.cells] for r in table.rows]
    if not rows:
        return ""
    head = "| " + " | ".join(rows[0]) + " |"
    sep = "|" + "|".join("---" for _ in rows[0]) + "|"
    return "\n".join([head, sep] + ["| " + " | ".join(r) + " |" for r in rows[1:]])


def save_webp(blob, path):
    """Convert any raster image to WebP (scaled to MAX_WIDTH). Returns (width, height) or None."""
    try:
        im = Image.open(io.BytesIO(blob))
        im.load()
    except Exception:
        return None  # vector formats (EMF/WMF) Pillow cannot read; listed for redrawing
    if im.mode not in ("RGB", "RGBA"):
        im = im.convert("RGBA" if "A" in im.getbands() or im.mode == "P" else "RGB")
    if im.width > MAX_WIDTH:
        im = im.resize((MAX_WIDTH, round(im.height * MAX_WIDTH / im.width)), Image.LANCZOS)
    im.save(path, "WEBP", quality=WEBP_QUALITY, method=6)
    return im.size


def load_review():
    """{(deck basename, file): row} from docs/image-review.csv, if it exists."""
    if not os.path.exists(REVIEW):
        return {}
    with open(REVIEW, encoding="utf8", newline="") as f:
        return {(r["deck"], r["file"]): r for r in csv.DictReader(f)}


def render_slides(pptx_path, slide_numbers, out_dir):
    """Export the given slides (1-based) to PNG with PowerPoint. Returns {n: png path}."""
    if not slide_numbers:
        return {}
    try:
        import win32com.client  # pywin32
    except ImportError:
        print("  (diagrams not rendered: pywin32 / PowerPoint unavailable)")
        return {}
    app = win32com.client.DispatchEx("PowerPoint.Application")
    pres = app.Presentations.Open(os.path.abspath(pptx_path), True, False, False)  # read-only, no window
    try:
        out = {}
        for n in slide_numbers:
            png = os.path.join(out_dir, f"_render-{n:02d}.png")
            pres.Slides(n).Export(png, "PNG", 1600, 900)
            out[n] = png
        return out
    finally:
        pres.Close()
        app.Quit()


def image_ref(name, kind, slide_no, review, deck, rel):
    row = review.get((deck, name))
    if row and row.get("decision") == "keep" and row.get("alt", "").strip():
        return f"![{escape_prose(row['alt'].strip())}]({rel}{name})"
    status = row["decision"] if row and row.get("decision") else "pending review"
    return f"{{/* Image {name} ({kind}, slide {slide_no}): {status}. */}}"


def convert(pptx_path, slug, render_diagrams):
    deck = os.path.basename(pptx_path)
    media_dir = os.path.join(OUT, *slug.split("/"))
    shutil.rmtree(media_dir, ignore_errors=True)
    os.makedirs(media_dir)
    review = load_review()
    depth = slug.count("/") + 2
    rel = "../" * depth + f"assets/{slug}/"

    prs = Presentation(pptx_path)
    slides, diagrams = [], []
    for n, slide in enumerate(prs.slides, 1):
        shapes = walk(slide.shapes)
        drawn = sum(1 for s in shapes if s.shape_type in DRAWN and not (s.has_text_frame and s.text_frame.text.strip()))
        if drawn >= DIAGRAM_SHAPES:
            diagrams.append(n)
        slides.append((n, slide, shapes))
    rendered = render_slides(pptx_path, diagrams, media_dir) if render_diagrams else {}

    seen, rows, sections = {}, [], []
    page_title = subtitle = None
    for n, slide, shapes in slides:
        title_shape = find_title(shapes)
        title_lines = [clean(p.text) for p in title_shape.text_frame.paragraphs] if title_shape is not None else []
        title_lines = [t for t in title_lines if t]
        title = re.sub(r"\s+", " ", title_lines[0]) if title_lines else None
        body = []
        for s in shapes:
            if s is title_shape or is_title(s):
                continue
            if s.has_text_frame and not re.search(r"[A-Za-z]{2}", s.text_frame.text):
                continue  # step numbers, emoji icons and other decoration
            if s.shape_type == MSO_SHAPE_TYPE.PICTURE or (s.is_placeholder and hasattr(s, "image")):
                try:
                    blob, ext = s.image.blob, s.image.ext
                except Exception:
                    continue
                key = hashlib.sha1(blob).hexdigest()[:10]
                if key not in seen:
                    name = f"s{n:02d}-{key}.webp"
                    size = save_webp(blob, os.path.join(media_dir, name))
                    if size is None:
                        name = f"s{n:02d}-{key}.{ext}"
                        with open(os.path.join(media_dir, name), "wb") as f:
                            f.write(blob)
                        size = ("", "")
                    seen[key] = name
                    rows.append({"deck": deck, "slide": n, "file": name, "kind": "picture",
                                 "width": size[0], "height": size[1], "source_name": s.name})
                    body.append(image_ref(name, "picture", n, review, deck, rel))
            elif s.has_table if hasattr(s, "has_table") else False:
                body.append(table_md(s.table))
            elif getattr(s, "has_chart", False) and s.has_chart:
                body.append(f"{{/* TODO: chart on slide {n} ({s.chart.chart_type}); rebuild as a figure or table. */}}")
            elif s.has_text_frame:
                md = text_frame_md(s.text_frame)
                if md:
                    body.append(md)
        if n in rendered:
            name = f"s{n:02d}-diagram.webp"
            with open(rendered[n], "rb") as f:
                size = save_webp(f.read(), os.path.join(media_dir, name))
            os.remove(rendered[n])
            rows.append({"deck": deck, "slide": n, "file": name, "kind": "diagram",
                         "width": size[0], "height": size[1], "source_name": "whole slide"})
            body.insert(0, image_ref(name, "diagram", n, review, deck, rel))
        elif n in diagrams:
            body.insert(0, f"{{/* TODO: slide {n} is a diagram of drawn shapes; render it (--render-diagrams) or redraw it. */}}")
        if slide.has_notes_slide:
            notes = clean(slide.notes_slide.notes_text_frame.text)
            if notes:
                body.append("\n\n".join(escape_prose(p.strip()) for p in notes.split("\n") if p.strip()))
        if n == 1 and page_title is None:
            page_title = title
            subtitle = " ".join(title_lines[1:]) or None  # the rest of the title box is a subtitle
            if subtitle and re.fullmatch(r"(bhutan\s+)?(art-trees\s+)?hfld readiness training", (title or "").lower()):
                page_title, subtitle = subtitle.split(":")[-1].strip(), None  # generic course title
            continue  # the title slide becomes the page title
        if title and SKIP_TITLES.match(title):
            continue
        sections.append((title, body, n))

    out, last = [], None
    for title, body, n in sections:
        if title and title != last:
            out.append(f"## {escape_prose(title)}")
        elif not title:
            out.append(f"{{/* Slide {n} has no title. */}}")
        last = title or last
        out.extend(b for b in body if b)

    with open(os.path.join(media_dir, "images.csv"), "w", encoding="utf8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["deck", "slide", "file", "kind", "width", "height", "source_name"])
        w.writeheader()
        w.writerows(rows)

    page_title = nice_title(page_title or slug.rsplit("/", 1)[-1].replace("-", " ").capitalize())
    description = quote_yaml(nice_title(subtitle)) if subtitle else "TODO"
    front = f"---\ntitle: {quote_yaml(page_title)}\ndescription: {description}\n---\n\n"
    mdx = front + f"{{/* Converted from {deck} by scripts/convert/pptx2mdx.py. Draft: edit before publishing. */}}\n\n" + "\n\n".join(out) + "\n"
    with open(os.path.join(OUT, slug + ".mdx"), "w", encoding="utf8", newline="\n") as f:
        f.write(mdx)
    return mdx, rows, media_dir, review, deck


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("pptx")
    ap.add_argument("slug", help="page path under src/content/docs, e.g. day-1/introduction-to-earth-engine")
    ap.add_argument("--render-diagrams", action="store_true", help="render diagram slides with PowerPoint (Windows)")
    ap.add_argument("--install-images", action="store_true", help="copy images marked keep to src/assets/<slug>/")
    args = ap.parse_args()
    if not re.fullmatch(r"[a-z0-9-]+(/[a-z0-9-]+)*", args.slug):
        sys.exit("slug must look like day-1/introduction-to-earth-engine")

    mdx, rows, media_dir, review, deck = convert(args.pptx, args.slug, args.render_diagrams)
    kept = [r["file"] for r in rows if review.get((deck, r["file"]), {}).get("decision") == "keep"]
    if args.install_images:
        dest = os.path.join(ROOT, "src", "assets", *args.slug.split("/"))
        shutil.rmtree(dest, ignore_errors=True)
        if kept:
            os.makedirs(dest)
            for name in kept:
                shutil.copy2(os.path.join(media_dir, name), dest)
    size = sum(os.path.getsize(os.path.join(media_dir, f)) for f in os.listdir(media_dir)) / 1e6
    print(f"{os.path.join(OUT, args.slug)}.mdx\n  {mdx.count(chr(10) + '## ')} sections, "
          f"{sum(r['kind'] == 'picture' for r in rows)} pictures, {sum(r['kind'] == 'diagram' for r in rows)} diagrams "
          f"({size:.1f} MB), {len(kept)} reviewed as keep, {mdx.count('```js')} code blocks, {mdx.count(chr(10) + '| ')} table rows")


if __name__ == "__main__":
    main()
